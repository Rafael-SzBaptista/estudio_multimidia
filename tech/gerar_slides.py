#!/usr/bin/env python3
"""
Gera slides PPTX a partir de arquivos de letra (.txt ou .md).

Uso (a partir da pasta raiz do projeto):
  python tech/gerar_slides.py
  python tech/gerar_slides.py letras/minha_musica.txt
  python tech/gerar_slides.py --fundo assets/backgrounds/natureza_01.jpg letras/

Ou simplesmente: gerar.bat
"""

from __future__ import annotations

import argparse
import hashlib
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

try:
    from PIL import ImageFont
except ImportError:
    ImageFont = None  # type: ignore

# ---------------------------------------------------------------------------
# Configuração
# ---------------------------------------------------------------------------

# Este arquivo fica em tech/; a raiz do projeto é a pasta pai.
TECH_DIR = Path(__file__).resolve().parent
BASE = TECH_DIR.parent
LETRAS_DIR = BASE / "letras"
OUTPUT_DIR = BASE / "output"
BACKGROUNDS_DIR = BASE / "assets" / "backgrounds"
DOCS_DIR = BASE / "README.md"

FONT_NAME = "Tw Cen MT"
FONT_SIZE_PT = 64.5
TITLE_FONT_SIZE_PT = 72
AUTHOR_FONT_SIZE_PT = 42
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
# Largura útil do texto (margens laterais de ~0,4")
USABLE_WIDTH_PT = (13.333 - 0.8) * 72

# Candidatos à fonte bold no Windows
FONT_CANDIDATES = [
    Path(r"C:\Windows\Fonts\TCB_____.TTF"),  # Tw Cen MT Bold
    Path(r"C:\Windows\Fonts\arialbd.ttf"),
    Path(r"C:\Windows\Fonts\ARIALBD.TTF"),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
]


# ---------------------------------------------------------------------------
# Leitura da letra
# ---------------------------------------------------------------------------

def parse_letra(path: Path) -> tuple[str, str, list[str]]:
    """
    Formato esperado:

        Título da música
        Nome do autor

        Linha da letra
        Outra linha
        ...

    Linhas em branco no meio da letra são ignoradas.
    Linhas começando com # são comentários.
    Aceita também "Título: ..." e "Autor: ..." nas duas primeiras linhas.
    """
    raw = path.read_text(encoding="utf-8-sig")
    lines = [ln.strip() for ln in raw.splitlines()]

    # Remove comentários e linhas vazias iniciais
    cleaned: list[str] = []
    for ln in lines:
        if ln.startswith("#"):
            continue
        cleaned.append(ln)

    # Pula vazios no começo
    while cleaned and not cleaned[0]:
        cleaned.pop(0)

    if len(cleaned) < 2:
        raise ValueError(
            f"{path.name}: o arquivo precisa ter pelo menos título, autor e a letra.\n"
            "Exemplo:\n  Título\n  Autor\n\n  linha 1\n  linha 2"
        )

    title = cleaned[0]
    author = cleaned[1]

    # Remove prefixos opcionais
    for prefix in ("título:", "titulo:", "title:"):
        if title.lower().startswith(prefix):
            title = title[len(prefix):].strip()
            break
    for prefix in ("autor:", "artista:", "author:"):
        if author.lower().startswith(prefix):
            author = author[len(prefix):].strip()
            break

    body_start = 2
    # Se a 3ª linha for vazia (separador), pula
    if body_start < len(cleaned) and not cleaned[body_start]:
        body_start += 1

    lyrics = [ln for ln in cleaned[body_start:] if ln]

    if not lyrics:
        raise ValueError(f"{path.name}: não encontrei linhas de letra após o autor.")

    return title, author, lyrics


# ---------------------------------------------------------------------------
# Medição e quebra de linhas
# ---------------------------------------------------------------------------

def find_font_path() -> Path | None:
    for p in FONT_CANDIDATES:
        if p.is_file():
            return p
    return None


def load_measure_font(size: float = FONT_SIZE_PT):
    """Fonte usada só para medir largura do texto."""
    if ImageFont is None:
        return None
    path = find_font_path()
    if path is None:
        return None
    # Pillow usa tamanho em px; com dpi 72, 1 pt ≈ 1 px
    return ImageFont.truetype(str(path), int(round(size)))


def text_width(font, text: str, size: float = FONT_SIZE_PT) -> float:
    if font is None:
        # Estimativa conservadora (~0,55 * size por caractere CAPS)
        return len(text) * size * 0.55
    bbox = font.getbbox(text)
    return float(bbox[2] - bbox[0])


def wrap_to_width(
    text: str, font, max_width: float, size: float = FONT_SIZE_PT
) -> list[str]:
    """Quebra uma linha em pedaços que cabem em max_width (por palavras)."""
    text = text.strip()
    if not text:
        return []
    if text_width(font, text, size) <= max_width:
        return [text]

    words = text.split()
    parts: list[str] = []
    current: list[str] = []

    for word in words:
        trial = " ".join(current + [word])
        if current and text_width(font, trial, size) > max_width:
            parts.append(" ".join(current))
            current = [word]
        else:
            current.append(word)

    if current:
        parts.append(" ".join(current))

    # Palavra isolada ainda maior que a largura: mantém mesmo assim
    return parts


def _lyric_slide(line1: str, line2: str | None = None) -> list[tuple[str, float]]:
    if line2 is None:
        return [(line1, FONT_SIZE_PT)]
    return [(line1, FONT_SIZE_PT), (line2, FONT_SIZE_PT)]


def build_slides(
    title: str, author: str, lyrics: list[str], font
) -> list[list[tuple[str, float]]]:
    """
    Monta lista de slides; cada slide é uma lista de (texto, tamanho_pt).
    Primeiro slide: título (72) + autor (42).
    Demais: letra em no máx. 2 linhas visuais.
    Texto em CAIXA ALTA.
    """
    slides: list[list[tuple[str, float]]] = []

    # --- Título + autor (primeiro slide) ---
    title_font = load_measure_font(TITLE_FONT_SIZE_PT)
    author_font = load_measure_font(AUTHOR_FONT_SIZE_PT)
    title_parts = wrap_to_width(
        title.upper(), title_font, USABLE_WIDTH_PT, TITLE_FONT_SIZE_PT
    )
    author_parts = wrap_to_width(
        author.upper(), author_font, USABLE_WIDTH_PT, AUTHOR_FONT_SIZE_PT
    )
    header = [(part, TITLE_FONT_SIZE_PT) for part in title_parts]
    header.extend((part, AUTHOR_FONT_SIZE_PT) for part in author_parts)
    if header:
        slides.append(header)

    # --- Letra ---
    # Empilha linhas curtas (1 visual) para formar pares.
    # Linhas longas (2+ visuais) viram slide(s) próprios, sem misturar com a seguinte.
    pending: str | None = None

    def flush_pending() -> None:
        nonlocal pending
        if pending is not None:
            slides.append(_lyric_slide(pending))
            pending = None

    for lyric_line in lyrics:
        parts = wrap_to_width(lyric_line.upper(), font, USABLE_WIDTH_PT, FONT_SIZE_PT)
        if not parts:
            continue

        if len(parts) >= 2:
            flush_pending()
            for i in range(0, len(parts), 2):
                a = parts[i]
                b = parts[i + 1] if i + 1 < len(parts) else None
                slides.append(_lyric_slide(a, b))
            continue

        # Uma linha visual
        if pending is None:
            pending = parts[0]
        else:
            slides.append(_lyric_slide(pending, parts[0]))
            pending = None

    flush_pending()
    return slides


# ---------------------------------------------------------------------------
# Fundo
# ---------------------------------------------------------------------------

def list_backgrounds() -> list[Path]:
    if not BACKGROUNDS_DIR.is_dir():
        return []
    exts = {".jpg", ".jpeg", ".png", ".webp"}
    return sorted(
        p for p in BACKGROUNDS_DIR.iterdir()
        if p.suffix.lower() in exts and p.is_file()
    )


def pick_background(title: str, override: Path | None) -> Path:
    if override is not None:
        if not override.is_file():
            raise FileNotFoundError(f"Imagem de fundo não encontrada: {override}")
        return override

    backgrounds = list_backgrounds()
    if not backgrounds:
        # Fallback: qualquer jpg em assets/
        assets = BASE / "assets"
        fallbacks = sorted(assets.glob("**/*.*")) if assets.is_dir() else []
        backgrounds = [
            p for p in fallbacks
            if p.suffix.lower() in {".jpg", ".jpeg", ".png", ".webp"}
        ]

    if not backgrounds:
        raise FileNotFoundError(
            "Nenhuma imagem de fundo encontrada.\n"
            f"Coloque arquivos .jpg/.png em: {BACKGROUNDS_DIR}"
        )

    # Escolha estável por título (mesmo fundo sempre para a mesma música)
    digest = hashlib.md5(title.encode("utf-8")).hexdigest()
    idx = int(digest, 16) % len(backgrounds)
    return backgrounds[idx]


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------

def set_run_font(run, size: float = FONT_SIZE_PT) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    rPr = run._r.get_or_add_rPr()
    for child in list(rPr):
        if child.tag.endswith("}latin") or child.tag.endswith("}cs") or child.tag.endswith("}ea"):
            rPr.remove(child)
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT_NAME)


def add_shadow(shape) -> None:
    spPr = shape._element.spPr
    effectLst = spPr.find(qn("a:effectLst"))
    if effectLst is None:
        effectLst = etree.SubElement(spPr, qn("a:effectLst"))
    outer = etree.SubElement(effectLst, qn("a:outerShdw"))
    outer.set("blurRad", "50800")
    outer.set("dist", "38100")
    outer.set("dir", "2700000")
    outer.set("algn", "ctr")
    outer.set("rotWithShape", "0")
    srgb = etree.SubElement(outer, qn("a:srgbClr"))
    srgb.set("val", "000000")
    alpha = etree.SubElement(srgb, qn("a:alpha"))
    alpha.set("val", "70000")


def add_slide(prs, blank, bg: Path, lines: list[tuple[str, float]]) -> None:
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(
        str(bg), Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height
    )

    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0.25),
        prs.slide_width,
        Inches(2.7),
    )
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(0, 0, 0)
    spPr = band._element.spPr
    solidFill = spPr.find(qn("a:solidFill"))
    srgb = solidFill.find(qn("a:srgbClr"))
    alpha = etree.SubElement(srgb, qn("a:alpha"))
    alpha.set("val", "40000")
    band.line.fill.background()

    txBox = slide.shapes.add_textbox(
        Inches(0.4),
        Inches(0.4),
        prs.slide_width - Inches(0.8),
        Inches(2.4),
    )
    tf = txBox.text_frame
    tf.word_wrap = False

    for i, (text, size) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(4)
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = text
        set_run_font(run, size)

    add_shadow(txBox)


def safe_filename(title: str) -> str:
    name = re.sub(r"[^\w\s\-À-ÿ]+", "", title, flags=re.UNICODE)
    name = re.sub(r"\s+", "_", name.strip())
    return name[:80] or "slides"


def create_pptx(
    title: str,
    author: str,
    lyrics: list[str],
    bg: Path,
    out_path: Path,
    font,
) -> int:
    slides = build_slides(title, author, lyrics, font)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    for lines in slides:
        add_slide(prs, blank, bg, lines)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return len(slides)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def collect_inputs(paths: list[Path]) -> list[Path]:
    files: list[Path] = []
    for p in paths:
        if p.is_dir():
            files.extend(sorted(p.glob("*.txt")))
            files.extend(sorted(p.glob("*.md")))
        elif p.is_file():
            files.append(p)
        else:
            print(f"Aviso: não encontrado: {p}", file=sys.stderr)
    # Remove duplicatas preservando ordem; ignora USO/README
    seen = set()
    result = []
    skip_names = {"uso.md", "readme.md", "script.md"}
    for f in files:
        key = f.resolve()
        if key in seen:
            continue
        if f.name.lower() in skip_names:
            continue
        seen.add(key)
        result.append(f)
    return result


def main(argv: list[str] | None = None) -> int:
    # Evita erro de encoding no terminal Windows (cp1252)
    if hasattr(sys.stdout, "reconfigure"):
        try:
            sys.stdout.reconfigure(encoding="utf-8", errors="replace")
            sys.stderr.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

    parser = argparse.ArgumentParser(
        description="Gera slides PPTX a partir de arquivos de letra (.txt / .md).",
        epilog="Sem argumentos, processa todos os .txt/.md da pasta letras/.",
    )
    parser.add_argument(
        "entradas",
        nargs="*",
        type=Path,
        help="Arquivo(s) ou pasta(s) com letras. Padrão: letras/",
    )
    parser.add_argument(
        "--fundo",
        type=Path,
        default=None,
        help="Caminho de uma imagem de fundo (usada em todos os slides do arquivo).",
    )
    parser.add_argument(
        "--saida",
        type=Path,
        default=OUTPUT_DIR,
        help=f"Pasta de saída (padrão: {OUTPUT_DIR})",
    )
    args = parser.parse_args(argv)

    entradas = args.entradas or [LETRAS_DIR]
    files = collect_inputs(entradas)

    if not files:
        print(
            "Nenhum arquivo de letra encontrado.\n"
            f"Coloque arquivos .txt na pasta: {LETRAS_DIR}\n"
            "Veja a pasta letras/ e o guia em README.md/USO.md.",
            file=sys.stderr,
        )
        return 1

    font = load_measure_font()
    if font is None:
        print(
            "Aviso: não foi possível carregar a fonte Tw Cen MT para medir texto.\n"
            "Usarei uma estimativa. No Windows, instale 'Tw Cen MT' para melhor resultado.",
            file=sys.stderr,
        )

    print(f"Arquivos a processar: {len(files)}\n")
    ok = 0
    for path in files:
        try:
            title, author, lyrics = parse_letra(path)
            bg = pick_background(title, args.fundo)
            out_name = safe_filename(title) + ".pptx"
            out_path = args.saida / out_name
            n = create_pptx(title, author, lyrics, bg, out_path, font)
            print(f"OK  {path.name}")
            print(f"    -> {out_path.name}  ({n} slides)")
            print(f"    fundo: {bg.name}")
            print(f"    {title} -- {author}\n")
            ok += 1
        except Exception as exc:
            print(f"ERRO  {path.name}: {exc}\n", file=sys.stderr)

    print(f"Concluido: {ok}/{len(files)} arquivo(s) gerado(s) em {args.saida}")
    return 0 if ok == len(files) else 2


if __name__ == "__main__":
    raise SystemExit(main())
