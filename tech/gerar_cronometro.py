#!/usr/bin/env python3
"""Gera cronômetro de 5 minutos (PPTX) compatível com PowerPoint e Google Slides."""

from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
from lxml import etree
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.util import Inches

try:
    import imageio.v2 as imageio
    import imageio_ffmpeg
except ImportError:
    imageio = None  # type: ignore
    imageio_ffmpeg = None  # type: ignore

TECH_DIR = Path(__file__).resolve().parent
BASE = TECH_DIR.parent
OUT_DIR = BASE / "cronômetro"
BG_PATH = OUT_DIR / "fundo_acampamento.jpg"
WIDE_PATH = OUT_DIR / "fundo_16x9.png"
PPTX_PATH = OUT_DIR / "cronometro_5min.pptx"
GIF_PATH = OUT_DIR / "cronometro_5min.gif"
SLIDES_PPTX_PATH = OUT_DIR / "cronometro_5min_slides.pptx"

W, H = 1920, 1080
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TOTAL_SEC = 5 * 60
HOLD_ZERO_SEC = 5

OVERLAY_BOX = (420, 470, 1500, 820)
TIME_CY = 640
TIME_FONT_PT = 200


def _font(name: str, size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(Path(r"C:\Windows\Fonts") / name), size)


def cover_resize(im: Image.Image, size: tuple[int, int]) -> Image.Image:
    tw, th = size
    scale = max(tw / im.width, th / im.height)
    nw, nh = int(im.width * scale + 0.5), int(im.height * scale + 0.5)
    im = im.resize((nw, nh), Image.Resampling.LANCZOS)
    left = (nw - tw) // 2
    top = (nh - th) // 2
    return im.crop((left, top, left + tw, top + th))


def build_base(photo: Image.Image) -> Image.Image:
    """Arte 16:9 em tela cheia, sem faixas laterais."""
    return cover_resize(photo.convert("RGB"), (W, H))


def _glow_time(
    overlay: Image.Image,
    text: str,
    cy: int,
    font: ImageFont.FreeTypeFont,
) -> None:
    """Números brancos com glow amarelo-ouro, no estilo da referência."""
    glow = Image.new("RGBA", overlay.size, (0, 0, 0, 0))
    g = ImageDraw.Draw(glow)
    cx = W / 2
    g.text(
        (cx, cy),
        text,
        font=font,
        fill=(255, 220, 120, 255),
        anchor="mm",
        stroke_width=18,
        stroke_fill=(255, 186, 60, 230),
    )
    overlay.alpha_composite(glow.filter(ImageFilter.GaussianBlur(22)))
    draw = ImageDraw.Draw(overlay)
    draw.text((cx, cy), text, font=font, fill=(255, 255, 255, 255), anchor="mm")


def render_frame(
    base_rgba: Image.Image,
    remaining: int,
    time_font: ImageFont.FreeTypeFont,
) -> Image.Image:
    overlay = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    mm, ss = divmod(max(remaining, 0), 60)
    _glow_time(overlay, f"{mm:02d}:{ss:02d}", TIME_CY, time_font)
    out = base_rgba.copy()
    out.alpha_composite(overlay)
    return out.convert("RGB")


def compose_base(photo_base: Image.Image) -> Image.Image:
    return photo_base.convert("RGBA")


def remaining_sequence() -> list[int]:
    return list(range(TOTAL_SEC, 0, -1)) + [0] * HOLD_ZERO_SEC


def _set_advance(slide) -> None:
    xml = (
        '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'advClick="0" advTm="1000"/>'
    )
    slide._element.append(etree.fromstring(xml))


def _px_box_to_emu(prs) -> tuple[int, int, int, int]:
    l, t, r, b = OVERLAY_BOX
    left = int(prs.slide_width * l / W)
    top = int(prs.slide_height * t / H)
    width = int(prs.slide_width * (r - l) / W)
    height = int(prs.slide_height * (b - t) / H)
    return left, top, width, height


def create_pptx(bg_jpeg: Path, overlay_dir: Path, overlays: list[str], out_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    left, top, width, height = _px_box_to_emu(prs)

    first = True
    for name in overlays:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(bg_jpeg), 0, 0, prs.slide_width, prs.slide_height)
        slide.shapes.add_picture(str(overlay_dir / name), left, top, width, height)
        _set_advance(slide)
        if first:
            slide.notes_slide.notes_text_frame.text = (
                "Google Slides: Apresentar → menu ⋮ (canto inferior) "
                "→ Avanço automático → 1 segundo.\n"
                "Alternativa: insira cronometro_5min.gif em um slide "
                "(Inserir → Imagem) e apresente — o GIF começa sozinho."
            )
            first = False

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))


def create_gif_pptx(gif_path: Path, out_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(gif_path), 0, 0, prs.slide_width, prs.slide_height)
    prs.save(str(out_path))


def encode_gif_from_mp4(mp4_path: Path, gif_path: Path) -> None:
    ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
    vf = (
        "split[s0][s1];"
        "[s0]palettegen=max_colors=128:stats_mode=diff[p];"
        "[s1][p]paletteuse=dither=none"
    )
    subprocess.run(
        [ffmpeg, "-y", "-i", str(mp4_path), "-vf", vf, "-loop", "-1", str(gif_path)],
        check=True,
    )


def save_preview(base: Image.Image, path: Path, remaining: int) -> None:
    time_font = _font("calibrib.ttf", TIME_FONT_PT)
    render_frame(compose_base(base), remaining, time_font).save(path, quality=95)


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        try:
            sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

    if not BG_PATH.is_file():
        print(f"Imagem de fundo não encontrada: {BG_PATH}", file=sys.stderr)
        return 1

    args = list(sys.argv[1:] if argv is None else argv)
    preview_only = "--preview" in args

    print("Compondo fundo do acampamento...")
    src = WIDE_PATH if WIDE_PATH.is_file() else BG_PATH
    photo = Image.open(src).convert("RGB")
    base = build_base(photo)

    if preview_only:
        for remaining, name in (
            (TOTAL_SEC, "preview_05m.jpg"),
            (150, "preview_02m30.jpg"),
            (20, "preview_00m20.jpg"),
        ):
            prev = OUT_DIR / name
            save_preview(base, prev, remaining)
            print(f"Prévia: {prev}")
        return 0

    time_font = _font("calibrib.ttf", TIME_FONT_PT)
    scene = compose_base(base)
    sequence = remaining_sequence()

    with tempfile.TemporaryDirectory(prefix="cronometro_") as tmp:
        tmp_dir = Path(tmp)
        overlay_dir = tmp_dir / "ovl"
        overlay_dir.mkdir()
        bg_jpeg = tmp_dir / "bg.jpg"
        scene.convert("RGB").save(bg_jpeg, quality=90, optimize=True)

        print(f"Renderizando {len(sequence)} frames do cronômetro...")
        overlays: list[str] = []
        box = OVERLAY_BOX
        mp4_path = tmp_dir / "cronometro.mp4"
        writer = None
        if imageio is not None:
            writer = imageio.get_writer(
                str(mp4_path),
                fps=1,
                codec="libx264",
                ffmpeg_params=["-pix_fmt", "yuv420p", "-crf", "22", "-tune", "stillimage"],
                macro_block_size=1,
            )
        try:
            for i, remaining in enumerate(sequence):
                frame = render_frame(scene, remaining, time_font)
                name = f"{i:04d}.jpg"
                frame.crop(box).save(overlay_dir / name, quality=90, optimize=True)
                overlays.append(name)
                if writer is not None:
                    small = frame.resize((1280, 720), Image.Resampling.LANCZOS)
                    writer.append_data(np.asarray(small))
                if i == 0 or (i + 1) % 30 == 0 or i + 1 == len(sequence):
                    print(f"    frame {i + 1}/{len(sequence)}  {remaining // 60:02d}:{remaining % 60:02d}")
        finally:
            if writer is not None:
                writer.close()

        print("Montando PPTX de um slide (GIF, Google Slides + PowerPoint)...")
        if imageio is not None and imageio_ffmpeg is not None and mp4_path.is_file():
            print("Gerando GIF para o Google Slides...")
            encode_gif_from_mp4(mp4_path, GIF_PATH)
            create_gif_pptx(GIF_PATH, PPTX_PATH)
            print(f"    GIF: {GIF_PATH}  ({GIF_PATH.stat().st_size / 1024 / 1024:.1f} MB)")

        print("Montando PPTX de slides (1 segundo cada, plano B)...")
        create_pptx(bg_jpeg, overlay_dir, overlays, SLIDES_PPTX_PATH)

    print(f"Pronto: {PPTX_PATH}")
    print("Este arquivo tem 1 slide com GIF — funciona no PowerPoint (F5) e no Google Slides.")
    print("Se o Google mostrar só 05:00 parado, use Inserir → Imagem → cronometro_5min.gif")
    print(f"Plano B (305 slides): {SLIDES_PPTX_PATH}")
    print("  No Google: Apresentar → ⋮ → Avanço automático → 1 segundo.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
